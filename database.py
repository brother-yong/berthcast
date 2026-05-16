import sqlite3
import os

DB_PATH = "berthai.db"


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    c = conn.cursor()

    c.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            org_name TEXT NOT NULL,
            model TEXT NOT NULL DEFAULT 'claude-haiku-4-5-20251001',
            is_admin INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS organisations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            model TEXT NOT NULL DEFAULT 'claude-haiku-4-5-20251001'
        );

        CREATE TABLE IF NOT EXISTS upload_sessions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            org_name TEXT NOT NULL,
            status TEXT DEFAULT 'uploading',
            context_json TEXT,
            dedup_confirmed INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS analysis_results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id INTEGER NOT NULL,
            inventory_report TEXT,
            recommendations_json TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (session_id) REFERENCES upload_sessions(id)
        );
    """)
    conn.commit()
    conn.close()


def excel_to_sqlite(filepath: str, table_name: str, session_id: int):
    """
    Convert an uploaded Excel file into a SQLite table scoped to the session.

    Synergix exports declare <dimension ref="A1"/> in their XML even for large files,
    which causes openpyxl read_only mode to return only 1 column per row.
    reset_dimensions() clears that bad metadata so all columns are read correctly.

    Metadata rows (printed by, date, company name, etc.) at the top are skipped
    automatically — the real header is the first row with 3+ non-empty cells.
    """
    import openpyxl

    conn = get_db()
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        ws.reset_dimensions()  # Fix for Synergix files with missing dimension declarations

        rows_iter = ws.iter_rows(values_only=True)

        # Skip metadata rows — find first row with 3+ non-empty values (the real header)
        raw_headers = None
        for row in rows_iter:
            non_empty = [v for v in row if v is not None and str(v).strip() != ""]
            if len(non_empty) >= 3:
                raw_headers = row
                break

        if raw_headers is None:
            return {"ok": False, "error": "Could not find data headers — file may be empty or unreadable."}

        # Sanitise column names
        seen = {}
        headers = []
        for i, h in enumerate(raw_headers):
            if h and str(h).strip():
                name = (str(h).strip().lower()
                        .replace(" ", "_").replace("/", "_")
                        .replace("-", "_").replace(".", "")
                        .replace("(", "").replace(")", ""))
            else:
                name = f"col_{i}"
            # Deduplicate
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 0
            headers.append(name)

        scoped_table = f"{table_name}_{session_id}"

        cols_def   = ", ".join(f'"{h}" TEXT' for h in headers) + ', "_session_id" TEXT'
        conn.execute(f'DROP TABLE IF EXISTS "{scoped_table}"')
        conn.execute(f'CREATE TABLE "{scoped_table}" ({cols_def})')

        placeholders = ", ".join("?" * (len(headers) + 1))
        insert_sql   = f'INSERT INTO "{scoped_table}" VALUES ({placeholders})'

        BATCH = 2000
        batch = []
        total = 0

        for row in rows_iter:
            values = [str(v) if v is not None else None for v in row[:len(headers)]]
            while len(values) < len(headers):
                values.append(None)
            values.append(str(session_id))
            batch.append(values)

            if len(batch) >= BATCH:
                conn.executemany(insert_sql, batch)
                conn.commit()
                total += len(batch)
                batch = []

        if batch:
            conn.executemany(insert_sql, batch)
            conn.commit()
            total += len(batch)

        wb.close()
        return {"ok": True, "rows": total, "table": scoped_table}

    except Exception as e:
        return {"ok": False, "error": str(e)}
    finally:
        conn.close()


def query(sql: str, params=()) -> list[dict]:
    conn = get_db()
    try:
        c = conn.cursor()
        c.execute(sql, params)
        rows = c.fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def execute(sql: str, params=()):
    conn = get_db()
    try:
        c = conn.cursor()
        c.execute(sql, params)
        conn.commit()
        return c.lastrowid
    finally:
        conn.close()


def table_exists(table_name: str) -> bool:
    rows = query(
        "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
        (table_name,)
    )
    return len(rows) > 0


def get_session_tables(session_id: int) -> dict:
    """Return which data tables exist for a given session."""
    expected = ["inventory", "purchase_orders", "sales", "suppliers", "customers", "stockouts"]
    return {
        name: table_exists(f"{name}_{session_id}")
        for name in expected
    }
